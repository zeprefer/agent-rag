from __future__ import annotations

from pathlib import Path
from typing import Iterable

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_LEFT
from reportlab.lib.pagesizes import LETTER
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import (
    KeepTogether,
    PageBreak,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)


ROOT = Path(__file__).resolve().parents[1]
OUTPUT_DIR = ROOT / "sample_data" / "chengyue_enterprise_kb" / "knowledge_base_files"
DOCX_PATH = OUTPUT_DIR / "05_travel_and_expense_policy.docx"
PDF_PATH = OUTPUT_DIR / "10_customer_support_sla_v2.2.pdf"

CJK_FONT = "Microsoft YaHei"
NAVY = RGBColor(11, 37, 69)
BLUE = RGBColor(46, 116, 181)
DARK_BLUE = RGBColor(31, 77, 120)
MUTED = RGBColor(90, 99, 110)
LIGHT_BLUE = "E8EEF5"
LIGHT_GRAY = "F2F4F7"


def set_run_font(run, size=None, bold=None, color=None, italic=None):
    run.font.name = CJK_FONT
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), CJK_FONT)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), CJK_FONT)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), CJK_FONT)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color is not None:
        run.font.color.rgb = color


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths_dxa: list[int]):
    total = sum(widths_dxa)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    for tag in ("w:tblW", "w:tblInd", "w:tblLayout"):
        existing = tbl_pr.find(qn(tag))
        if existing is not None:
            tbl_pr.remove(existing)
    tbl_w = OxmlElement("w:tblW")
    tbl_w.set(qn("w:w"), str(total))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_pr.append(tbl_w)
    tbl_ind = OxmlElement("w:tblInd")
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")
    tbl_pr.append(tbl_ind)
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tbl_pr.append(layout)

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        for index, cell in enumerate(row.cells):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(widths_dxa[index]))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def add_numbering_definition(doc: Document, num_id: int, abstract_id: int, ordered: bool):
    numbering = doc.part.numbering_part.element
    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    multi = OxmlElement("w:multiLevelType")
    multi.set(qn("w:val"), "singleLevel")
    abstract.append(multi)
    lvl = OxmlElement("w:lvl")
    lvl.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:start")
    start.set(qn("w:val"), "1")
    lvl.append(start)
    num_fmt = OxmlElement("w:numFmt")
    num_fmt.set(qn("w:val"), "decimal" if ordered else "bullet")
    lvl.append(num_fmt)
    lvl_text = OxmlElement("w:lvlText")
    lvl_text.set(qn("w:val"), "%1." if ordered else "•")
    lvl.append(lvl_text)
    suff = OxmlElement("w:suff")
    suff.set(qn("w:val"), "tab")
    lvl.append(suff)
    p_pr = OxmlElement("w:pPr")
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "num")
    tab.set(qn("w:pos"), "540")
    tabs.append(tab)
    p_pr.append(tabs)
    ind = OxmlElement("w:ind")
    ind.set(qn("w:left"), "540")
    ind.set(qn("w:hanging"), "270")
    p_pr.append(ind)
    spacing = OxmlElement("w:spacing")
    spacing.set(qn("w:after"), "80")
    spacing.set(qn("w:line"), "300")
    spacing.set(qn("w:lineRule"), "auto")
    p_pr.append(spacing)
    lvl.append(p_pr)
    numbering.append(abstract)
    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abstract_ref = OxmlElement("w:abstractNumId")
    abstract_ref.set(qn("w:val"), str(abstract_id))
    num.append(abstract_ref)
    numbering.append(num)


def set_num(paragraph, num_id):
    p_pr = paragraph._p.get_or_add_pPr()
    num_pr = p_pr.find(qn("w:numPr"))
    if num_pr is None:
        num_pr = OxmlElement("w:numPr")
        p_pr.append(num_pr)
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    num = OxmlElement("w:numId")
    num.set(qn("w:val"), str(num_id))
    num_pr.append(ilvl)
    num_pr.append(num)


def set_docx_styles(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = CJK_FONT
    normal._element.rPr.rFonts.set(qn("w:ascii"), CJK_FONT)
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), CJK_FONT)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), CJK_FONT)
    normal.font.size = Pt(11)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    specs = {
        "Heading 1": (16, BLUE, 18, 10),
        "Heading 2": (13, BLUE, 14, 7),
        "Heading 3": (12, DARK_BLUE, 10, 5),
    }
    for name, (size, color, before, after) in specs.items():
        style = styles[name]
        style.font.name = CJK_FONT
        style._element.rPr.rFonts.set(qn("w:ascii"), CJK_FONT)
        style._element.rPr.rFonts.set(qn("w:hAnsi"), CJK_FONT)
        style._element.rPr.rFonts.set(qn("w:eastAsia"), CJK_FONT)
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = color
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    add_numbering_definition(doc, num_id=41, abstract_id=41, ordered=False)
    add_numbering_definition(doc, num_id=42, abstract_id=42, ordered=True)


def add_docx_title(doc):
    section = doc.sections[0]
    header = section.header
    p = header.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run("澄岳智造 | 财务运营制度")
    set_run_font(run, 8.5, True, MUTED)

    footer = section.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = fp.add_run("CY-FIN-012  |  第 ")
    set_run_font(r, 8.5, False, MUTED)
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    fp._p.append(fld)
    r = fp.add_run(" 页")
    set_run_font(r, 8.5, False, MUTED)

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run("差旅与费用报销管理办法")
    set_run_font(r, 27, True, NAVY)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(14)
    r = p.add_run("可执行政策参考 | 版本 3.2 | 2026-03-01 生效")
    set_run_font(r, 11.5, False, MUTED)

    rows = [
        ("文档编号", "CY-FIN-012", "所有者", "财务运营部"),
        ("密级", "L1 内部公开", "替代版本", "3.1"),
        ("适用对象", "全体员工及经批准的外部人员", "咨询队列", "FIN-EXPENSE"),
    ]
    table = doc.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    for row_data in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row_data):
            cells[i].text = value
    table._tbl.remove(table.rows[0]._tr)
    set_table_geometry(table, [1350, 3330, 1350, 3330])
    tr_pr = table.rows[0]._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    tr_pr.append(repeat)
    for row in table.rows:
        for i, cell in enumerate(row.cells):
            if i in (0, 2):
                set_cell_shading(cell, LIGHT_BLUE)
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(0)
                for run in p.runs:
                    set_run_font(run, 9.3, i in (0, 2), NAVY if i in (0, 2) else None)
    doc.add_paragraph()


def add_body(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.widow_control = True
    r = p.add_run(text)
    set_run_font(r, 11)
    return p


def add_bullet(doc, text):
    p = doc.add_paragraph()
    set_num(p, 41)
    r = p.add_run(text)
    set_run_font(r, 11)
    return p


def add_numbered(doc, text):
    p = doc.add_paragraph()
    set_num(p, 42)
    r = p.add_run(text)
    set_run_font(r, 11)
    return p


def add_policy_table(doc, headers, rows, widths):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for i, header in enumerate(headers):
        table.rows[0].cells[i].text = header
        set_cell_shading(table.rows[0].cells[i], LIGHT_BLUE)
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = value
    set_table_geometry(table, widths)
    tr_pr = table.rows[0]._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    tr_pr.append(repeat)
    for row_index, row in enumerate(table.rows):
        for col_index, cell in enumerate(row.cells):
            for p in cell.paragraphs:
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.space_after = Pt(2)
                p.paragraph_format.line_spacing = 1.15
                if col_index > 0 and len(cell.text) < 14:
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for run in p.runs:
                    set_run_font(run, 9.2, row_index == 0, NAVY if row_index == 0 else None)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def build_docx():
    doc = Document()
    set_docx_styles(doc)
    add_docx_title(doc)

    doc.add_heading("1. 总则", level=1)
    add_body(doc, "本办法规范因公差旅、客户招待、日常业务费用和公司卡使用。费用必须真实、必要、合理、与公司业务相关，并在发生前取得适用审批。预算存在不代表可以绕过审批、票据或采购规则。")
    add_body(doc, "员工不得通过拆分行程、拆分发票、改变费用类别或让同事代报规避限额。审批人对业务必要性和成本中心负责，财务对票据、政策与会计处理负责；财务审核不能替代业务审批。")

    doc.add_heading("2. 差旅申请与预订", level=1)
    add_body(doc, "跨城市出行应在预订前提交差旅申请，包含目的、地点、日期、项目码、预算和客户信息。普通差旅至少提前 3 个工作日申请；国际差旅至少提前 10 个工作日。ORION-27 相关差旅必须使用财务项目码 CY-CS-2026-027。")
    for item in [
        "使用公司指定平台预订机票、火车和酒店；平台无合适资源时，保留比价证据后可自行预订。",
        "行程改变导致退改费时，应在报销说明中写明原因；个人原因产生的额外费用由个人承担。",
        "结合多个业务目的的行程，应按主要受益项目分摊；不能确定时由成本中心负责人确认。",
        "因 P1 或 SEV-1/2 紧急出行无法事前申请的，应由事件指挥官口头授权，并在 1 个工作日内补录。",
    ]:
        add_bullet(doc, item)

    doc.add_heading("3. 交通标准", level=1)
    add_policy_table(
        doc,
        ["交通类型", "默认标准", "可升级条件"],
        [
            ("国内航班", "经济舱；优先合理低价直达方案", "连续飞行超过 6 小时且分管副总裁事前批准可选公务舱"),
            ("国际航班", "经济舱", "单段计划飞行超过 8 小时，经分管副总裁批准可选公务舱"),
            ("高铁/动车", "二等座", "单程超过 4 小时可选一等座；商务座需 CFO 例外批准"),
            ("市内交通", "公共交通或合规网约车", "携带重物、深夜安全、多人同行更经济或客户现场无公共交通"),
            ("自驾私车", "事前批准，按 1.2 元/公里补贴", "补贴已包含燃油和折旧；停车、路桥费可凭票另报"),
        ],
        [1800, 3000, 4560],
    )
    add_body(doc, "航空公司会员升舱券可以使用，但公司只承担政策允许舱位价格。员工不得为了积累积分选择明显更贵或不合理的行程。个人里程、积分和会员权益归员工，因业务退票形成的现金退款必须退回公司。")

    doc.add_heading("4. 住宿标准", level=1)
    add_body(doc, "住宿限额按每间每晚含税价计算。公司指定会议酒店、客户指定安全酒店或大型展会期间无法取得限额内住宿时，可在预订前由部门负责人批准上浮，报销时附价格证据。")
    add_policy_table(
        doc,
        ["城市类别", "示例", "经理及以下", "总监及以上"],
        [
            ("A 类", "北京、上海、深圳、新加坡", "650 元或等值当地货币", "850 元或等值当地货币"),
            ("B 类", "苏州、广州、杭州、成都、南京", "500 元", "650 元"),
            ("C 类", "其他城市", "380 元", "500 元"),
        ],
        [1350, 2700, 2655, 2655],
    )
    add_body(doc, "深圳代码为 CN-SHEN，苏州为 CN-SZ。员工在深圳出差应适用 A 类标准，不能因缩写 SZ 误用苏州标准。共享房间不作强制要求。住宿中包含的早餐不再重复领取对应餐费额度。")

    doc.add_heading("5. 餐费与杂费", level=1)
    add_body(doc, "国内差旅按实际发生、限额内报销，不采用无票固定津贴。单人每日餐费上限：A 类城市 220 元、B 类 180 元、C 类 150 元。每日限额可在早餐、午餐和晚餐之间调剂，但不得跨天结转。")
    add_body(doc, "由会议、酒店、客户或同事报销提供的餐食应从当日上限扣减：早餐 20%，午餐 35%，晚餐 45%。例如 A 类城市酒店含早餐，当日员工自费午晚餐的剩余额度为 176 元。若仅一顿晚餐花费 220 元而当日其他餐均由员工自理，仍可在 220 元日上限内报销；报销的是每日总额，不是单餐上限。")
    add_body(doc, "洗衣费仅在连续出差达到 7 晚时可报，每满 7 晚上限 100 元。迷你吧、付费影视、健身、个人护理、罚款、违章和陪同人员费用不得报销。")

    doc.add_heading("6. 客户招待与礼品", level=1)
    add_body(doc, "客户招待必须有明确业务目的、参加人姓名与组织、日期、地点和讨论主题。员工人数原则上不得超过外部参加人数。人均含税 500 元以内由部门负责人事前批准；超过 500 元或单次总额超过 5,000 元由分管副总裁批准。")
    add_body(doc, "政府相关方、招标评审人员、正在处理索赔的客户或对方政策禁止的情形，不得提供礼品或招待。现金、购物卡、充值卡和可兑换现金的礼物一律禁止。公司品牌纪念品单件不超过 200 元时仍需符合对方政策。")

    doc.add_heading("7. 可报销与不可报销项目", level=1)
    for item in [
        "可报销：经批准的交通住宿、签证与必要保险、客户现场防护用品、业务通信、合理行李费、必要打印和翻译。",
        "需额外说明：临时购买软件、居家办公设备、团队活动、个人手机国际漫游和紧急替代设备。",
        "不可报销：通勤、个人旅行、家庭成员费用、罚款、酒吧娱乐、个人会员、非业务礼品、遗失物品和未授权升级。",
    ]:
        add_bullet(doc, item)
    add_body(doc, "单笔软件或服务费用达到采购门槛时，不能通过报销流程替代采购。即使低于 10,000 元，只要供应商访问 L3/L4 数据、企业身份或生产环境，仍须完成采购和安全评审。")

    doc.add_heading("8. 票据、时限与汇率", level=1)
    add_body(doc, "员工应在行程结束或费用发生后 10 个工作日内提交。超过 30 个自然日需要部门负责人说明；超过 90 个自然日原则上不予报销，除非存在出院、长期现场封闭、系统故障或法律要求等可验证原因，并通过 FIN-EX-09。")
    add_body(doc, "中国大陆费用应提供合规发票；境外费用提供当地有效收据。电子票据必须清晰、完整且不得重复提交。原始票据遗失时，员工提交付款证据和遗失说明，单笔不超过 500 元由部门负责人批准；超过 500 元还需财务负责人批准。")
    add_body(doc, "外币由系统按交易日公司卡汇率或费用发生日的财务月度汇率折算。员工个人银行卡收取的外汇手续费可凭证报销，但汇率差损益不另行补偿。")

    doc.add_heading("9. 公司卡", level=1)
    add_body(doc, "公司卡优先用于机票、酒店和批准的高频业务支出。持卡人不得转借卡片或共享验证码。误刷个人消费应在 2 个工作日内报告，并在对账单到期前返还。丢失或可疑交易应立即冻结并通知发卡行和财务。")

    doc.add_heading("10. 审批流程", level=1)
    for item in [
        "员工提交费用，确认业务目的、项目码、参加人、票据和例外说明。",
        "直属经理确认必要性、真实性和交付关系；项目经理确认项目费用归属。",
        "成本中心负责人确认预算与授权；超过门槛或特殊类别进入追加审批。",
        "财务检查政策、票据、税务、重复报销和会计科目，退回时说明具体缺口。",
        "批准后进入付款批次；通常在完整批准后的 7 个工作日内支付。",
    ]:
        add_numbered(doc, item)

    doc.add_heading("11. 例外与抽查", level=1)
    add_body(doc, "财务例外使用 FIN-EX-09，必须列出条款、金额、原因、补偿性控制和到期日。批准例外不等于自动批准费用，员工仍需完成正常报销。财务可抽查路线、参加人、价格和重复票据；故意虚报、篡改票据或拆分费用将转交人力与合规处理。")

    doc.add_heading("12. 示例", level=1)
    add_policy_table(
        doc,
        ["场景", "处理结论"],
        [
            ("经理在深圳住 620 元/晚", "深圳为 A 类，经理及以下上限 650 元；在其他条件满足时可报。"),
            ("成都酒店含早餐，当日另报 170 元餐费", "B 类每日 180 元扣早餐 20% 后余额 144 元，超出 26 元需个人承担或事前例外。"),
            ("单程 5 小时高铁选择一等座", "单程超过 4 小时可选一等座，无需因舱位单独申请例外。"),
            ("8,000 元 SaaS 可接触客户生产数据", "金额低于报价竞争门槛，但仍必须走采购与信息安全评审，不能直接费用报销。"),
        ],
        [3100, 6260],
    )

    doc.add_heading("13. 版本说明", level=1)
    add_body(doc, "3.2 版将深圳明确归入 A 类城市，统一外币折算规则，新增含餐扣减比例和软件服务安全评审提示。本版自生效日起替代 3.1 版。")
    doc.core_properties.title = "澄岳智造差旅与费用报销管理办法"
    doc.core_properties.subject = "企业知识库测试样例"
    doc.core_properties.author = "澄岳智造（虚构测试数据）"
    doc.core_properties.keywords = "差旅, 报销, 费用, 酒店, 餐费"
    doc.save(DOCX_PATH)


def register_pdf_font():
    candidates = [
        Path("C:/Windows/Fonts/msyh.ttc"),
        Path("C:/Windows/Fonts/msyh.ttf"),
        Path("C:/Windows/Fonts/simhei.ttf"),
        Path("C:/Windows/Fonts/simsun.ttc"),
    ]
    for candidate in candidates:
        if candidate.exists():
            pdfmetrics.registerFont(TTFont("KB-CJK", str(candidate)))
            return "KB-CJK"
    raise RuntimeError("No CJK TrueType font found")


def pdf_header_footer(canvas, doc):
    canvas.saveState()
    canvas.setFont("KB-CJK", 8.5)
    canvas.setFillColor(colors.HexColor("#667085"))
    canvas.drawString(0.8 * inch, 10.45 * inch, "澄岳智造 | 客户支持服务等级说明")
    canvas.drawRightString(7.7 * inch, 0.48 * inch, f"CY-SUP-005  |  第 {doc.page} 页")
    canvas.restoreState()


def build_pdf():
    font = register_pdf_font()
    styles = getSampleStyleSheet()
    title = ParagraphStyle("KBTitle", parent=styles["Title"], fontName=font, fontSize=25, leading=31, textColor=colors.HexColor("#0B2545"), alignment=TA_LEFT, spaceAfter=8)
    subtitle = ParagraphStyle("KBSubtitle", parent=styles["Normal"], fontName=font, fontSize=11, leading=16, textColor=colors.HexColor("#667085"), spaceAfter=16)
    h1 = ParagraphStyle("KBH1", parent=styles["Heading1"], fontName=font, fontSize=16, leading=21, textColor=colors.HexColor("#2E74B5"), spaceBefore=16, spaceAfter=8, keepWithNext=True)
    h2 = ParagraphStyle("KBH2", parent=styles["Heading2"], fontName=font, fontSize=12.5, leading=17, textColor=colors.HexColor("#1F4D78"), spaceBefore=10, spaceAfter=6, keepWithNext=True)
    body = ParagraphStyle("KBBody", parent=styles["BodyText"], fontName=font, fontSize=10.5, leading=16, textColor=colors.HexColor("#20242A"), alignment=TA_LEFT, spaceAfter=7)
    small = ParagraphStyle("KBSmall", parent=body, fontSize=8.6, leading=12, spaceAfter=0)
    callout = ParagraphStyle("KBCallout", parent=body, fontSize=10, leading=15, leftIndent=8, rightIndent=8, borderPadding=8, backColor=colors.HexColor("#F4F6F9"), spaceBefore=5, spaceAfter=10)

    doc = SimpleDocTemplate(str(PDF_PATH), pagesize=LETTER, leftMargin=0.8 * inch, rightMargin=0.8 * inch, topMargin=0.78 * inch, bottomMargin=0.72 * inch, title="客户支持服务等级说明 2.2", author="澄岳智造（虚构测试数据）")
    story = []
    story.append(Spacer(1, 0.25 * inch))
    story.append(Paragraph("客户支持服务等级说明", title))
    story.append(Paragraph("版本 2.2 | 2026-03-01 生效 | L2 受控", subtitle))
    meta = Table([
        [Paragraph("文档编号", small), Paragraph("CY-SUP-005", small), Paragraph("所有者", small), Paragraph("全球技术支持部", small)],
        [Paragraph("替代版本", small), Paragraph("2.1", small), Paragraph("适用范围", small), Paragraph("AtlasOps、CY-Edge 与 SenseHub", small)],
        [Paragraph("状态", small), Paragraph("现行", small), Paragraph("解释优先级", small), Paragraph("客户合同优先于本通用说明", small)],
    ], colWidths=[0.95 * inch, 2.05 * inch, 1.05 * inch, 2.85 * inch])
    meta.setStyle(TableStyle([
        ("FONTNAME", (0, 0), (-1, -1), font), ("FONTSIZE", (0, 0), (-1, -1), 8.6),
        ("BACKGROUND", (0, 0), (0, -1), colors.HexColor("#E8EEF5")),
        ("BACKGROUND", (2, 0), (2, -1), colors.HexColor("#E8EEF5")),
        ("TEXTCOLOR", (0, 0), (-1, -1), colors.HexColor("#20242A")),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#C8D1DC")),
        ("LEFTPADDING", (0, 0), (-1, -1), 7), ("RIGHTPADDING", (0, 0), (-1, -1), 7),
        ("TOPPADDING", (0, 0), (-1, -1), 6), ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
    ]))
    story.append(meta)
    story.append(Spacer(1, 12))
    story.append(Paragraph("重要：支持计划规定响应目标；产品订阅层级、硬件保修和专业服务不自动提供相同权益。发生冲突时，以客户已签署订单和合同为准。", callout))

    sections = [
        ("1. 服务范围", [
            "本说明适用于已购买有效支持服务的客户，涵盖受支持版本中的产品使用问题、故障诊断、事件协调和知识指导。它不替代实施项目、定制开发、客户网络运维、第三方产品支持或硬件保修。",
            "支持服务目标是响应与协作目标，不是所有事件在目标时间内永久修复的保证。恢复可以通过回滚、绕行、容量调整、隔离或临时配置实现；根因修复可能在恢复后交付。",
        ]),
        ("2. 支持计划", []),
    ]
    for heading, paras in sections:
        story.append(Paragraph(heading, h1))
        for para in paras:
            story.append(Paragraph(para, body))

    plan_data = [
        ["计划", "覆盖时间", "P1 初次响应", "P2 初次响应", "主要权益"],
        ["标准支持", "合同区域业务日 09:00-18:00", "60 分钟", "4 个业务小时", "在线工单、知识库、P3 1 个业务日"],
        ["企业支持", "P1/P2 7x24；其余业务时间", "30 分钟", "2 小时", "电话升级、季度支持回顾"],
        ["关键业务支持", "P1/P2 7x24；其余业务时间", "10 分钟", "1 小时", "P1 恢复目标 4 小时、指定事件协调"],
    ]
    plan_table = Table([[Paragraph(str(cell), small) for cell in row] for row in plan_data], colWidths=[0.9 * inch, 1.65 * inch, 1.05 * inch, 1.05 * inch, 2.25 * inch], repeatRows=1)
    plan_table.setStyle(TableStyle([
        ("FONTNAME", (0, 0), (-1, -1), font), ("FONTSIZE", (0, 0), (-1, -1), 8.3),
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#E8EEF5")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.HexColor("#0B2545")),
        ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#C8D1DC")),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("ALIGN", (2, 1), (3, -1), "CENTER"),
        ("LEFTPADDING", (0, 0), (-1, -1), 5), ("RIGHTPADDING", (0, 0), (-1, -1), 5),
        ("TOPPADDING", (0, 0), (-1, -1), 6), ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
    ]))
    story.append(plan_table)
    story.append(Paragraph("“Enterprise 产品订阅”和“企业支持计划”是不同商品。购买 AtlasOps Enterprise 不代表自动享有企业支持。", callout))

    story.append(Paragraph("3. 事件优先级", h1))
    priority_rows = [
        ["级别", "判断条件", "不属于该级别的例子"],
        ["P1 严重", "生产服务大范围不可用；核心业务停止且无合理绕行；确认重大数据完整性风险；或远程控制产生迫近安全风险。", "单个测试设备、一般报表延迟、功能咨询、未发布功能请求。"],
        ["P2 高", "重要生产功能严重退化但存在绕行，或影响范围受限且可能扩大。", "无业务影响的轻微界面缺陷。"],
        ["P3 一般", "一般缺陷、配置问题、性能咨询或不紧急的使用问题。", "需要立即恢复的大范围生产中断。"],
        ["P4 请求", "文档建议、最佳实践、功能建议和计划性咨询。", "已有生产影响的缺陷。"],
    ]
    t = Table([[Paragraph(cell, small) for cell in row] for row in priority_rows], colWidths=[0.85 * inch, 3.7 * inch, 2.35 * inch], repeatRows=1)
    t.setStyle(TableStyle([
        ("FONTNAME", (0, 0), (-1, -1), font), ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#E8EEF5")),
        ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#C8D1DC")), ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("LEFTPADDING", (0, 0), (-1, -1), 5), ("RIGHTPADDING", (0, 0), (-1, -1), 5),
        ("TOPPADDING", (0, 0), (-1, -1), 5), ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
    ]))
    story.append(t)
    story.append(Paragraph("优先级由实际业务影响、范围和绕行方案决定，不由客户职级、客户健康颜色或情绪单独决定。支持可在获得新证据后调整优先级。", body))

    story.append(Paragraph("4. 计时规则", h1))
    for para in [
        "初次响应从支持系统确认收到工单时开始。初次响应应由具备处理能力的人员确认影响、请求必要信息并说明下一步；自动回执不单独算作初次响应。",
        "P1 恢复目标从 P1 正式声明时开始。等待客户提供必需日志、远程访问、复现条件或批准操作期间，恢复计时可以暂停；暂停必须在工单中写明所缺信息和恢复计时条件。初次响应计时不因等待客户而暂停。",
        "若事件从 P2 升为 P1，P1 初次响应以升级时间开始，已发生的处理历史仍应保留。若支持建议从 P1 降级，需向客户联系人说明依据；关键业务支持的 P1 降级应取得客户确认，客户 30 分钟未回应且影响已解除时可临时降级并继续记录。",
        "业务小时按合同指定支持区域的工作日和节假日计算。中国大陆与新加坡虽然同为 UTC+8，业务日日历可能不同。",
    ]:
        story.append(Paragraph(para, body))

    story.append(Paragraph("5. P1 事件协作", h1))
    steps = [
        "创建事件桥接并确认事件指挥官、技术负责人和客户沟通负责人。",
        "记录已知影响、开始时间、受影响版本与区域、近期变更和安全信号。",
        "关键业务支持至少每 30 分钟更新一次客户；企业支持至少每 60 分钟更新一次，除非合同另有约定。",
        "恢复后验证业务路径、积压、数据完整性和监控，再结束 P1。",
        "符合重大事件条件时，通常在恢复后 5 个工作日内提供复盘初稿。",
    ]
    for idx, step in enumerate(steps, 1):
        story.append(Paragraph(f"{idx}. {step}", body))

    story.append(Paragraph("6. 客户责任", h1))
    for para in [
        "客户应维护授权联系人，提供准确影响、设备或租户标识、时间、版本、近期变更和可安全共享的诊断信息，并配合验证。客户不得在工单中粘贴密码、私钥或生产根凭据。",
        "客户负责其网络、运营商、供电、第三方系统、自研代码和不受支持版本。支持团队会协助定位边界，但相关恢复时间不计入澄岳可控制的服务指标，除非合同另有规定。",
    ]:
        story.append(Paragraph(para, body))

    story.append(Paragraph("7. 排除与限制", h1))
    exclusions = [
        "计划维护且已按合同通知；",
        "客户或第三方未经批准的配置、代码、网络或硬件导致的影响；",
        "客户拒绝实施合理修复、绕行或受支持版本升级；",
        "不可抗力及超出合理控制的公共基础设施故障，但团队仍会提供可行协助；",
        "预览、试用、实验室或已停止支持的版本，除非书面约定。",
    ]
    for item in exclusions:
        story.append(Paragraph("• " + item, body))

    story.append(Paragraph("8. 服务抵扣", h1))
    for para in [
        "服务抵扣只在合同明确约定时适用，不自动支付。客户应在受影响月份结束后 30 个自然日内提交申请，列明租户、事件、时间和影响。逾期申请可被拒绝。",
        "若合同采用本通用计算，月可用性低于 99.9% 但不低于 99.5% 时，抵扣当月受影响服务费的 5%；低于 99.5% 但不低于 99.0% 时为 10%；低于 99.0% 时为 20%。单月总抵扣不超过当月受影响服务费的 20%，且抵扣是可用性的唯一通用财务补救。",
        "初次响应或恢复目标未达到不必然产生抵扣，除非合同明确把该目标与抵扣关联。硬件保修、专业服务费、税费和第三方费用不计入抵扣基数。",
    ]:
        story.append(Paragraph(para, body))

    story.append(Paragraph("9. 事件复盘", h1))
    story.append(Paragraph("复盘应包含影响、时间线、根因、促成因素、恢复动作和预防措施。复盘初稿中的行动项可能仍为开放状态；只有达到验证标准后才能标记完成。安全事件的对外披露范围由信息安全、法务和合同所有者决定。", body))

    story.append(Paragraph("10. 联系与升级", h1))
    story.append(Paragraph("常规问题通过支持门户提交。企业支持和关键业务支持的 P1 可使用订单中列明的 7x24 电话。疑似账号失陷、数据误发或密钥暴露应同时使用安全事件通道；公司内部员工使用 SEC-INCIDENT 或分机 7711。媒体询问不得由支持人员直接回应。", body))

    story.append(Paragraph("11. 与内部连续性目标的关系", h1))
    story.append(Paragraph("AtlasOps 内部连续性目标 BCP-RTO-042 为重大区域故障下目标恢复时间 4 小时、目标数据点 15 分钟。该目标用于内部设计与验证，不自动扩大客户合同权利。关键业务支持的 P1 恢复目标同为 4 小时，但二者的适用条件、计时和法律效果不同。", body))

    story.append(Paragraph("12. 版本说明", h1))
    story.append(Paragraph("2.2 版自 2026-03-01 生效并替代 2.1 版。关键业务支持 P1 初次响应从 15 分钟缩短为 10 分钟，恢复目标从 6 小时缩短为 4 小时；新增 P1 降级确认、服务抵扣通用计算和内部连续性目标边界。历史事件按当时生效版本评估。", callout))

    doc.build(story, onFirstPage=pdf_header_footer, onLaterPages=pdf_header_footer)


def main():
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    build_docx()
    build_pdf()
    print(DOCX_PATH)
    print(PDF_PATH)


if __name__ == "__main__":
    main()
