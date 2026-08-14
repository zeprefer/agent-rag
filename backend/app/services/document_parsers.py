from dataclasses import dataclass
from io import BytesIO
from pathlib import Path
import re
import subprocess
import tempfile

from docx import Document as DocxDocument
from pypdf import PdfReader

from app.core.config import settings


@dataclass(frozen=True)
class ParsedSection:
    text: str
    heading_path: str | None = None
    page_number: int | None = None


class DocumentParseError(RuntimeError):
    pass


def parse_document_bytes(*, file_type: str, filename: str, data: bytes) -> list[ParsedSection]:
    if file_type == "txt":
        return _parse_text(data)
    if file_type == "pdf":
        return _parse_pdf(data)
    if file_type == "docx":
        return _parse_docx(data)
    if file_type == "doc":
        return _parse_doc(data=data, filename=filename)
    if file_type == "md":
        return _parse_markdown(data)
    raise DocumentParseError(f"Unsupported parser for file type: {file_type}")


def _parse_text(data: bytes) -> list[ParsedSection]:
    text = data.decode("utf-8-sig", errors="replace")
    cleaned = _clean_text(text)
    if not cleaned:
        raise DocumentParseError("No extractable text found in text file")
    return [ParsedSection(text=cleaned)]


def _parse_pdf(data: bytes) -> list[ParsedSection]:
    try:
        reader = PdfReader(BytesIO(data))
    except Exception as exc:
        raise DocumentParseError(f"PDF could not be opened: {exc}") from exc

    sections: list[ParsedSection] = []
    for page_index, page in enumerate(reader.pages, start=1):
        text = page.extract_text() or ""
        text = _clean_text(text)
        if text:
            sections.append(ParsedSection(text=text, page_number=page_index))
    if not sections:
        raise DocumentParseError("No extractable text found in PDF")
    return sections


def _parse_docx(data: bytes) -> list[ParsedSection]:
    try:
        document = DocxDocument(BytesIO(data))
    except Exception as exc:
        raise DocumentParseError(f"DOCX could not be opened: {exc}") from exc

    sections: list[ParsedSection] = []
    headings: list[str] = []
    buffer: list[str] = []

    def flush() -> None:
        text = _clean_text("\n".join(buffer))
        if text:
            sections.append(ParsedSection(text=text, heading_path=" > ".join(headings) or None))
        buffer.clear()

    for paragraph in document.paragraphs:
        text = _clean_text(paragraph.text)
        if not text:
            continue

        style_name = paragraph.style.name if paragraph.style is not None else ""
        heading_level = _heading_level(style_name)
        if heading_level is not None:
            flush()
            headings[:] = headings[: heading_level - 1]
            headings.append(text)
        else:
            buffer.append(text)

    for table in document.tables:
        rows: list[str] = []
        for row in table.rows:
            cells = [_clean_text(cell.text) for cell in row.cells]
            row_text = " | ".join(cell for cell in cells if cell)
            if row_text:
                rows.append(row_text)
        if rows:
            buffer.append("\n".join(rows))

    flush()
    if not sections:
        raise DocumentParseError("No extractable text found in DOCX")
    return sections


def _parse_doc(*, data: bytes, filename: str) -> list[ParsedSection]:
    with tempfile.TemporaryDirectory() as temp_dir:
        input_path = Path(temp_dir) / filename
        input_path.write_bytes(data)

        command = [
            settings.libreoffice_binary,
            "--headless",
            "--convert-to",
            "docx",
            "--outdir",
            temp_dir,
            str(input_path),
        ]
        try:
            result = subprocess.run(command, capture_output=True, text=True, timeout=90, check=False)
        except FileNotFoundError as exc:
            raise DocumentParseError("LibreOffice binary not found; cannot parse .doc files") from exc
        except subprocess.TimeoutExpired as exc:
            raise DocumentParseError("DOC conversion timed out") from exc

        if result.returncode != 0:
            detail = result.stderr.strip() or result.stdout.strip()
            raise DocumentParseError(f"DOC conversion failed: {detail}")

        converted_files = list(Path(temp_dir).glob("*.docx"))
        if not converted_files:
            raise DocumentParseError("DOC conversion did not produce a DOCX file")

        return _parse_docx(converted_files[0].read_bytes())


def _parse_markdown(data: bytes) -> list[ParsedSection]:
    text = data.decode("utf-8-sig", errors="replace")
    sections: list[ParsedSection] = []
    headings: list[str] = []
    buffer: list[str] = []

    def flush() -> None:
        body = _clean_text("\n".join(buffer))
        if body:
            sections.append(ParsedSection(text=body, heading_path=" > ".join(headings) or None))
        buffer.clear()

    for line in text.splitlines():
        heading_match = re.match(r"^(#{1,6})\s+(.+)$", line.strip())
        if heading_match:
            flush()
            level = len(heading_match.group(1))
            title = _clean_text(heading_match.group(2))
            headings[:] = headings[: level - 1]
            headings.append(title)
            continue
        buffer.append(line)

    flush()
    if not sections:
        cleaned = _clean_text(text)
        if cleaned:
            sections.append(ParsedSection(text=cleaned))
    if not sections:
        raise DocumentParseError("No extractable text found in Markdown")
    return sections


def _heading_level(style_name: str) -> int | None:
    match = re.match(r"Heading\s+([1-6])", style_name)
    if not match:
        return None
    return int(match.group(1))


def _clean_text(text: str) -> str:
    text = text.replace("\x00", "")
    lines = [re.sub(r"[ \t]+", " ", line).strip() for line in text.splitlines()]
    return "\n".join(line for line in lines if line).strip()
